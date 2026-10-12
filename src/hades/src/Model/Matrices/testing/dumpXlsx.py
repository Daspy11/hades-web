"""Dump an .xlsx workbook as JSON with openpyxl (used by Excel.test.ts).

Usage: python3 -I dumpXlsx.py <file.xlsx>
"""
import datetime
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def flag(x):
    return None if x is None else bool(x)


def colour(c):
    if c is None:
        return None
    return {"rgb": c.rgb if isinstance(c.rgb, str) else None, "auto": flag(c.auto)}


def side(s):
    if s is None or s.style is None:
        return None
    return {"style": s.style, "color": colour(s.color)}


def value(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return {"date": v.isoformat()}
    return v


def main(path):
    wb = openpyxl.load_workbook(path)
    out = {"sheets": wb.sheetnames, "data": {}}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                cells[c.coordinate] = {
                    "value": value(c.value),
                    "type": c.data_type,
                    "bold": bool(c.font.b),
                    "fontName": c.font.name,
                    "fontColor": colour(c.font.color),
                    "fill": c.fill.fill_type,
                    "fillFg": colour(c.fill.fgColor) if c.fill.fill_type else None,
                    "fillBg": colour(c.fill.bgColor) if c.fill.fill_type else None,
                    "border": {k: side(getattr(c.border, k)) for k in ("left", "right", "top", "bottom")},
                    "numberFormat": c.number_format,
                    "horizontal": c.alignment.horizontal,
                    "vertical": c.alignment.vertical,
                    "wrap": flag(c.alignment.wrap_text),
                }
        widths = {}
        for i in range(1, ws.max_column + 1):
            letter = get_column_letter(i)
            if letter in ws.column_dimensions:
                d = ws.column_dimensions[letter]
                widths[letter] = {"width": d.width, "bestFit": flag(d.bestFit), "customWidth": flag(d.customWidth)}
        heights = {str(k): d.height for k, d in ws.row_dimensions.items() if d.height is not None}
        out["data"][ws.title] = {
            "cells": cells,
            "merges": sorted(str(r) for r in ws.merged_cells.ranges),
            "widths": widths,
            "heights": heights,
            "maxRow": ws.max_row,
            "maxColumn": ws.max_column,
        }
    json.dump(out, sys.stdout)


if __name__ == "__main__":
    main(sys.argv[1])
